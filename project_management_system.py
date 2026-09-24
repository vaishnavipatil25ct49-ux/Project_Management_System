import tkinter as tk
from tkinter import ttk,messagebox
from openpyxl import Workbook,load_workbook
import os

F="project_management.xlsx"

if not os.path.exists(F):
    w=Workbook();w.active.append(["ID","Name","Client","Start","Deadline","Status"]);w.save(F)

def clear():
    [e.delete(0,tk.END) for e in es];st.set("Not Started")

def show():
    t.delete(*t.get_children())
    for r in load_workbook(F).active.iter_rows(min_row=2,values_only=True):t.insert("",tk.END,values=r)

def add():
    d=[e.get() for e in es]+[st.get()]
    if not d[0] or not d[1]:return messagebox.showwarning("Warning","Enter ID and Name")
    w=load_workbook(F);w.active.append(d);w.save(F);clear();show()

def update():
    w=load_workbook(F);s=w.active;pid=es[0].get()
    for r in range(2,s.max_row+1):
        if str(s.cell(r,1).value)==pid:
            d=[e.get() for e in es]+[st.get()]
            for c in range(1,7):s.cell(r,c).value=d[c-1]
            w.save(F);show();return
    messagebox.showerror("Error","ID not found")

def delete():
    if not messagebox.askyesno("Delete","Delete record?"):return
    w=load_workbook(F);s=w.active;pid=es[0].get()
    for r in range(2,s.max_row+1):
        if str(s.cell(r,1).value)==pid:
            s.delete_rows(r);w.save(F);clear();show();return

def search():
    pid=searchbox.get();t.delete(*t.get_children())
    for r in load_workbook(F).active.iter_rows(min_row=2,values_only=True):
        if str(r[0])==pid:t.insert("",tk.END,values=r)

def select(e):
    d=t.item(t.focus())["values"]
    if d:
        clear()
        for x,v in zip(es,d[:5]):x.insert(0,v)
        st.set(d[5])

def dashboard():
    global es,st,t,searchbox,root
    root=tk.Tk();root.title("Project Management System");root.geometry("750x550")
    tk.Label(root,text="PROJECT MANAGEMENT SYSTEM",font=("Arial",14,"bold")).pack()
    
    f=tk.Frame(root);f.pack()
    es=[]
    for i,n in enumerate(["ID","Project Name","Client","Start Date","Deadline"]):
        tk.Label(f,text=n).grid(row=i,column=0)
        e=tk.Entry(f,width=25);e.grid(row=i,column=1);es.append(e)
    
    tk.Label(f,text="Status").grid(row=5)
    st=ttk.Combobox(f,values=["Not Started","In Progress","Completed"],state="readonly",width=22)
    st.set("Not Started");st.grid(row=5,column=1)

    b=tk.Frame(root);b.pack(pady=8)
    for x,cmd in [("Add",add),("Update",update),("Delete",delete),("Clear",clear)]:
        tk.Button(b,text=x,width=10,command=cmd).pack(side="left",padx=2)

    q=tk.Frame(root);q.pack()
    searchbox=tk.Entry(q,width=15);searchbox.pack(side="left")
    tk.Button(q,text="Search",command=search).pack(side="left")
    tk.Button(q,text="Show All",command=show).pack(side="left")

    t=ttk.Treeview(root,columns=("ID","Name","Client","Start","Deadline","Status"),show="headings")
    for c in t["columns"]:t.heading(c,text=c)
    t.pack(fill="both",expand=True);t.bind("<ButtonRelease-1>",select)
    show();root.mainloop()

def login():
    if u.get()=="admin" and p.get()=="1234":
        loginwin.destroy();dashboard()
    else:messagebox.showerror("Error","Wrong Login")

loginwin=tk.Tk();loginwin.title("Login");loginwin.geometry("300x200")
tk.Label(loginwin,text="PROJECT MANAGEMENT",font=("Arial",13,"bold")).pack(pady=15)
tk.Label(loginwin,text="Username").pack();u=tk.Entry(loginwin);u.pack()
tk.Label(loginwin,text="Password").pack();p=tk.Entry(loginwin,show="*");p.pack()
tk.Button(loginwin,text="LOGIN",command=login).pack(pady=15)
loginwin.mainloop()